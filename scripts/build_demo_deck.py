"""Build docs/memoria-demo.pptx, the deck for the Memoria demo.

Uses screenshots from web/screenshots/live/ (a real run against the live
API) and falls back to the Playwright suite's screenshots when a live one
is missing. Re-run after taking new screenshots:
    backend/.venv/bin/python scripts/build_demo_deck.py
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "web" / "screenshots"
OUT = ROOT / "docs" / "memoria-demo.pptx"

BG = RGBColor(0xF5, 0xF1, 0xE8)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x67, 0x60)
ACCENT = RGBColor(0x0F, 0x5C, 0x5C)
RULE = RGBColor(0xD9, 0xD3, 0xC7)
PANEL = RGBColor(0xFB, 0xF9, 0xF4)
FONT = "Helvetica Neue"
MONO = "Menlo"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
MARGIN = Inches(0.7)


def shot(live_name: str, fallback: str) -> Path:
    live = SHOTS / "live" / live_name
    return live if live.exists() else SHOTS / fallback


def text(slide, x, y, w, h, value, size=18, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT, spacing=1.15):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        run = p.add_run()
        run.text = line
        run.font.size, run.font.bold, run.font.name = Pt(size), bold, font
        run.font.color.rgb = color
    return box


def rule(slide, y, x=MARGIN, w=None):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x + (w or prs.slide_width - 2 * MARGIN), y)
    line.line.color.rgb = INK
    line.line.width = Pt(1.25)


def new_slide(kicker: str, title: str, notes: str):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    text(s, MARGIN, Inches(0.45), Inches(8), Inches(0.3), kicker.upper(), 11, ACCENT, True)
    text(s, MARGIN, Inches(0.75), Inches(11.9), Inches(0.9), title, 30, INK, True)
    rule(s, Inches(1.55))
    s.notes_slide.notes_text_frame.text = notes
    return s


def picture(slide, path: Path, x, y, w, h, anchor="top"):
    """Fit an image into the box, cropping (never stretching) to the box's aspect ratio."""
    with Image.open(path) as im:
        iw, ih = im.size
    box_ratio, img_ratio = w / h, iw / ih
    pic = slide.shapes.add_picture(str(path), x, y, w, h)
    if img_ratio < box_ratio:  # image taller than box: crop bottom (keep the top of the page)
        keep = img_ratio / box_ratio
        if anchor == "top":
            pic.crop_bottom = 1 - keep
        elif anchor == "bottom":
            pic.crop_top = 1 - keep
        else:
            pic.crop_top = pic.crop_bottom = (1 - keep) / 2
    else:
        keep = box_ratio / img_ratio
        pic.crop_left = pic.crop_right = (1 - keep) / 2
    pic.line.color.rgb = RULE
    pic.line.width = Pt(0.75)
    return pic


def card(slide, x, y, w, h, title, body, mono_title=False):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    rect.fill.solid()
    rect.fill.fore_color.rgb = PANEL
    rect.line.color.rgb = RULE
    rect.line.width = Pt(0.75)
    rect.shadow.inherit = False
    text(slide, x + Inches(0.2), y + Inches(0.18), w - Inches(0.4), Inches(0.4), title, 15, ACCENT if mono_title else INK,
         True, MONO if mono_title else FONT)
    text(slide, x + Inches(0.2), y + Inches(0.62), w - Inches(0.4), h - Inches(0.75), body, 13, MUTED)


# 1. Title
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = BG
text(s, MARGIN, Inches(2.2), Inches(11), Inches(0.4), "SALES ENGINEER TRACK  ·  CARD D", 12, ACCENT, True)
text(s, MARGIN, Inches(2.65), Inches(11), Inches(1.2), "Memoria", 66, INK, True)
rule(s, Inches(3.95), w=Inches(5))
text(s, MARGIN, Inches(4.2), Inches(10), Inches(1.2),
     "A voice-first sales engineer that interviews the rep, researches the deal,\n"
     "builds a verified campaign, and remembers every account.", 22, MUTED)
text(s, MARGIN, Inches(6.5), Inches(11), Inches(0.4), "Any company  ·  any product  ·  github.ibm.com/hazem-awadalla/memoria", 12, MUTED)
s.notes_slide.notes_text_frame.text = (
    "Memoria is the sales engineer track build. One line: it interviews the rep by voice, does the research, "
    "builds the campaign, checks every claim, and gets sharper on each account because it remembers.")

# 2. The problem
s = new_slide("The problem", "Final-pitch prep is scattered, and AI tools forget",
              "Three pains. Prep lives in heads and call notes. Generic AI writes plausible copy with invented numbers. "
              "And every chat starts from zero, so the new objection from last call never makes it into the plan.")
for i, (title, body) in enumerate([
    ("Context lives in the rep's head", "Stack, stakeholders, objections and competitive moves sit in call notes and memory, not in one place."),
    ("Generic AI invents specifics", "Plausible statistics, customer names and benchmarks with no source. One wrong number costs the deal's credibility."),
    ("Every session starts from zero", "The objection raised on Tuesday's call is gone by Thursday's prep. The pitch order never adapts."),
]):
    card(s, MARGIN + i * Inches(4.05), Inches(2.0), Inches(3.8), Inches(2.6), title, body)
text(s, MARGIN, Inches(5.1), Inches(11.9), Inches(1),
     "Card D test question, asked every session:  “Tomorrow’s call is the final pitch. What’s our strategy?”", 18, INK)

# 3. What it does
s = new_slide("What Memoria does", "Interview, research, personas, campaign, verify, remember",
              "Walk the pipeline left to right. The interview is the only part the rep has to sit through; everything after "
              "runs in the background while the screen shows each stage.")
steps = [("Interview", "Voice discovery,\nMEDDPICC slots,\none question per turn"),
         ("Research", "Uploaded docs and\nweb research become\na numbered evidence ledger"),
         ("Personas", "Buying committee,\nevery statement evidenced\nor labelled a hypothesis"),
         ("Campaign", "Positioning, objection\nmatrix, battle card,\npitch order, outreach"),
         ("Verify", "Every claim checked\nagainst the ledger;\none revision, then human review"),
         ("Remember", "Per-tenant memory store;\nobjections superseded,\nnever deleted")]
w = Inches(1.88)
for i, (title, body) in enumerate(steps):
    x = MARGIN + i * (w + Inches(0.12))
    text(s, x, Inches(2.0), w, Inches(0.3), f"0{i + 1}", 12, ACCENT, True, MONO)
    rule(s, Inches(2.4), x=x, w=w - Inches(0.1))
    text(s, x, Inches(2.55), w, Inches(0.4), title, 18, INK, True)
    text(s, x, Inches(3.05), w - Inches(0.1), Inches(1.6), body, 13, MUTED)
text(s, MARGIN, Inches(5.2), Inches(11.9), Inches(1.2),
     "Assets out: slide decks (.pptx), documents and whitepapers (.docx + .pdf), blogs (.md + .html), email sequences (.eml).\n"
     "Every fact carries an evidence marker; missing proof becomes a visible placeholder, never an invention.", 15, INK)

# 4. Interview screenshot
s = new_slide("Live demo  ·  the interview", "Memoria interviews the rep by voice",
              "Speech-to-text is Whisper running on WebGPU inside the browser, so audio never leaves the laptop. Memoria's "
              "voice is Kokoro, served from a Podman container, streamed sentence by sentence. Typing works too. "
              "Every answer lands in the slot board on the right with the rep's exact words as evidence.")
picture(s, shot("03-live-interview-mid.png", "interview-mid-conversation-1440.png"), MARGIN, Inches(1.85), Inches(7.6), Inches(5.2))
x = MARGIN + Inches(7.95)
for i, (title, body) in enumerate([
    ("Whisper on WebGPU", "Speech-to-text in the browser. Audio stays on the device."),
    ("Kokoro voice", "Text-to-speech in a Podman container, streamed per sentence. 0.54 s for a first sentence."),
    ("One question per turn", "Priority order: meeting, pain, metrics, decision, buyer, objections, competition."),
]):
    text(s, x, Inches(1.95) + i * Inches(1.65), Inches(4.0), Inches(0.4), title, 16, INK, True)
    text(s, x, Inches(2.35) + i * Inches(1.65), Inches(4.0), Inches(1.1), body, 13, MUTED)

# 5. Slot board
s = new_slide("Live demo  ·  discovery", "Every answer becomes evidence on the slot board",
              "Sixteen slots, grouped Deal, Pain, Decision, Risk, Campaign. Status is confirmed, unsure, open or conflict. "
              "Conflict fires when the rep contradicts memory, and Memoria asks which is current instead of guessing.")
picture(s, shot("04-live-slots.png", "slot-board-partial-1440.png"), MARGIN, Inches(1.85), Inches(11.9), Inches(5.3))

# 6. Architecture
s = new_slide("Architecture", "Three zones: browser, Podman, Anthropic",
              "Three zones. Browser: Next.js plus Whisper on WebGPU. Podman: the web, api and tts containers. Anthropic: "
              "Claude Sonnet 5.5 on the Messages API for the chain, with structured outputs and caching, web search for "
              "research, and Managed Agents with one memory store per tenant for long-term memory.")
zones = [("Browser", ["Next.js UI", "Whisper on WebGPU", "Audio playback queue"]),
         ("Podman", ["web  ·  Next.js :3000", "api  ·  FastAPI :8000", "tts  ·  Kokoro :8880"]),
         ("Anthropic", ["Messages API  ·  claude-sonnet-5-5", "web_search / web_fetch", "Managed Agents + memory stores"])]
zw = Inches(3.7)
for i, (title, items) in enumerate(zones):
    zx = MARGIN + i * (zw + Inches(0.4))
    text(s, zx, Inches(1.95), zw, Inches(0.4), title.upper(), 12, ACCENT, True)
    for j, item in enumerate(items):
        y = Inches(2.4) + j * Inches(1.05)
        rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, zx, y, zw, Inches(0.8))
        rect.fill.solid()
        rect.fill.fore_color.rgb = PANEL
        rect.line.color.rgb = RULE
        rect.shadow.inherit = False
        text(s, zx + Inches(0.2), y + Inches(0.24), zw - Inches(0.4), Inches(0.4), item, 14, INK, False,
             MONO if ":" in item or "_" in item else FONT)
    if i < 2:
        ax = zx + zw + Inches(0.05)
        arrow = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ax, Inches(3.85), ax + Inches(0.3), Inches(3.85))
        arrow.line.color.rgb = INK
        arrow.line.width = Pt(1.5)
text(s, MARGIN, Inches(5.85), Inches(11.9), Inches(1),
     "Two lanes: the interview is interactive and streamed; the campaign build runs in the background.\n"
     "The model writes structured JSON; FastAPI renders the files deterministically, so every claim is checked before rendering.", 14, MUTED)

# 7. Evidence or nothing
s = new_slide("Trust", "Evidence or nothing",
              "This is the part that makes it safe to send. Retrieval and research produce a numbered ledger, E1, E2 and so on, "
              "with a verbatim quote, source and date. The code assigns the IDs, not the model. Every asset must cite them. "
              "A strict verifier checks each claim; unsupported claims get one revision, then the asset ships flagged for human review.")
for i, (title, body) in enumerate([
    ("E# ledger", "Every fact: one-sentence claim, verbatim quote, source, date. IDs assigned by code, not trusted from the model."),
    ("[PLACEHOLDER: ...]", "Where proof is missing, the asset says so visibly. No invented statistics, customers, quotes or benchmarks."),
    ("Verifier gate", "Each claim is supported, unsupported, contradicted or opinion. Pass only if every claim holds."),
    ("One revision, then a human", "Needs revision triggers one reviser pass. A second failure ships with a needs-human-review verdict."),
]):
    card(s, MARGIN + (i % 2) * Inches(6.05), Inches(1.95) + (i // 2) * Inches(2.45), Inches(5.85), Inches(2.25), title, body, True)

# 8. Build complete
s = new_slide("Live demo  ·  the campaign", "Nine stages, five verified assets, ready to download",
              "Stage timeline on the right with durations. Each asset shows its verdict, how many placeholders the rep still "
              "needs to fill, and download links for every format.")
picture(s, shot("06-live-build-complete.png", "build-complete-assets-1440.png"), MARGIN, Inches(1.85), Inches(11.9), Inches(5.3))

# 9. Memory across sessions
s = new_slide("Memory  ·  Card D", "Session two is sharper because Memoria remembers",
              "Round one: stack overview, earlier objections, the pitch deck. Round two: a new objection from the latest call "
              "and a competitive update from product management. Same question both times. In session two the answer leads "
              "with what changed and the pitch order moves. The memory diff shows exactly what Memoria chose to remember.")
picture(s, shot("07-live-memory.png", "strategy-and-memory-diff-1440.png"), MARGIN, Inches(1.85), Inches(7.2), Inches(5.3), anchor="bottom")
x = MARGIN + Inches(7.55)
for i, (title, body) in enumerate([
    ("Round 1", "Customer stack, earlier objections, current pitch deck."),
    ("Round 2", "New objection from the latest call, plus a competitive update."),
    ("Session 2 answer", "Opens “Since our last session: ...”, answers the new objection, reorders the pitch."),
    ("One store per tenant", "Isolation is tested: a canary planted in one tenant must never appear in another’s output."),
]):
    text(s, x, Inches(1.95) + i * Inches(1.28), Inches(4.3), Inches(0.35), title, 15, INK, True)
    text(s, x, Inches(2.3) + i * Inches(1.28), Inches(4.3), Inches(0.85), body, 13, MUTED)

# 10. Built test-first
s = new_slide("Engineering", "Built test-first, measured before tuned",
              "Built by seven parallel tracks, each test-driven, then integrated. Numbers are from the final run. The eval "
              "follows the design constitution: multiple trials, rate plus pass@k and pass^k, a holdout tenant we never tune on.")
stats = [("505", "backend tests passing"), ("120", "web unit tests"), ("23", "Playwright end-to-end tests"), ("3", "fictional tenants, one holdout")]
for i, (n, label) in enumerate(stats):
    x = MARGIN + i * Inches(3.0)
    text(s, x, Inches(2.0), Inches(2.8), Inches(1.0), n, 54, ACCENT, True)
    text(s, x, Inches(3.05), Inches(2.8), Inches(0.6), label, 14, MUTED)
rule(s, Inches(4.0))
text(s, MARGIN, Inches(4.25), Inches(11.9), Inches(2.6),
     "Eval: a simulated rep plays a hidden fact sheet in speech-to-text style; deterministic checks for slot recall, the new objection,\n"
     "the competitive update, pitch-order change, superseded objections and tenant isolation; one-criterion LLM judges for the rest.\n"
     "Fake mode swaps every external service for deterministic fakes, so the full stack runs in CI with no API calls.", 15, INK)

# 11. Status
s = new_slide("Status", "What is proven, and what is next",
              "Be straight with the room. Everything on the left was run today. The right column is what we have not yet "
              "measured, and the eval harness exists to measure it.")
text(s, MARGIN, Inches(1.95), Inches(5.8), Inches(0.4), "PROVEN TODAY", 12, ACCENT, True)
text(s, MARGIN, Inches(2.4), Inches(5.8), Inches(4.2),
     "Live interview on Claude Sonnet 5.5 through the real UI\n"
     "Kokoro voice in Podman, through the API\n"
     "Full pipeline end to end in fake mode, UI and CLI\n"
     "All five asset types rendered and downloadable\n"
     "Session-two answer leads with what changed (fake mode)", 15, INK, spacing=1.5)
x = MARGIN + Inches(6.1)
text(s, x, Inches(1.95), Inches(5.8), Inches(0.4), "NEXT", 12, ACCENT, True)
text(s, x, Inches(2.4), Inches(5.8), Inches(4.2),
     "Baseline eval in live mode: 5 trials per task, holdout tenant\n"
     "Session two on live Managed Agents memory\n"
     "Whisper on WebGPU on presenter hardware\n"
     "Full podman compose stack in one command\n"
     "Rendered-asset visual review; footnotes in .docx", 15, INK, spacing=1.5)

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print(f"wrote {OUT.relative_to(ROOT)} ({len(prs.slides)} slides)")
for name in ["03-live-interview-mid.png", "04-live-slots.png", "06-live-build-complete.png", "07-live-memory.png"]:
    print(("live   " if (SHOTS / "live" / name).exists() else "suite  ") + name)
