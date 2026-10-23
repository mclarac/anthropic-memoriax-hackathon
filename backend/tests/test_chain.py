import asyncio
import base64
import io
from pathlib import Path

import pytest
from conftest import (
    TODAY,
    asset_content,
    brief,
    evidence,
    plan_entry,
    verdict,
)
from docx import Document as DocxDocument

from memoria.chain import BuildRunner
from memoria.chain_prompts import ASSET_SCHEMAS, ASSET_SYSTEMS, CHAIN, TRACK_TEST_QUESTION
from memoria.models import (
    AssetReady,
    BuildComplete,
    MemoryDiffReady,
    StageEvent,
    StrategyAnswerReady,
    Turn,
)
from memoria.storage import Storage

STAGE_ORDER = ["deal_brief", "retrieval", "research", "personas", "strategy",
               "assets", "verifier", "memory_writer", "strategy_answer"]


class Rig:
    def __init__(self, tmp_path: Path, doubles, profile, with_store: bool = True):
        self.doubles = doubles
        self.storage = Storage(tmp_path)
        self.tenant = self.storage.create_tenant(profile)
        if with_store:
            self.tenant.memory_store_id = "store-1"
            self.storage.save_tenant(self.tenant)
        self.session = self.storage.create_session(self.tenant.id, "globex", "dana")
        self.session.turns = [Turn(speaker="memoria", text="Hi"), Turn(speaker="rep", text="Hello")]
        self.runner = BuildRunner(doubles, self.storage, lambda: TODAY)

    def add_document(self, filename, data, media_type, scope="tenant", account_slug=None):
        ref = self.storage.add_document(self.tenant.id, filename=filename, media_type=media_type,
                                        data=data, scope=scope, account_slug=account_slug)
        self.tenant = self.storage.get_tenant(self.tenant.id)
        return ref

    async def run(self):
        return [e async for e in self.runner.run(self.session, self.tenant)]


@pytest.fixture
def rig(tmp_path, doubles, profile, render_recorder):
    r = Rig(tmp_path, doubles, profile)
    r.add_document("facts.md", b"# Facts\nWe are fast.", "text/markdown")
    return r


def stage_events(events):
    return [(e.stage.stage_id, e.stage.status) for e in events if isinstance(e, StageEvent)]


async def test_happy_path_emits_stages_in_chain_order_and_ends_with_build_complete(rig):
    events = await rig.run()

    expected = []
    for stage in STAGE_ORDER:
        expected += [(stage, "running"), (stage, "done")]
    assert stage_events(events) == expected
    assert isinstance(events[-1], BuildComplete)
    assert sum(isinstance(e, BuildComplete) for e in events) == 1
    assert rig.session.phase == "complete"
    assert rig.storage.get_session(rig.session.id).phase == "complete"


async def test_done_events_carry_duration_and_running_events_do_not(rig):
    events = await rig.run()
    for e in events:
        if isinstance(e, StageEvent):
            assert (e.stage.duration_ms is None) == (e.stage.status == "running")
            if e.stage.status == "done":
                assert e.stage.duration_ms >= 0


async def test_asset_ready_arrives_during_verifier_and_memory_events_before_build_complete(rig):
    events = await rig.run()
    kinds = [(type(e).__name__, getattr(getattr(e, "stage", None), "stage_id", None),
              getattr(getattr(e, "stage", None), "status", None)) for e in events]
    asset_at = next(i for i, e in enumerate(events) if isinstance(e, AssetReady))
    assert kinds.index(("StageEvent", "verifier", "running")) < asset_at < kinds.index(
        ("StageEvent", "verifier", "done"))
    names = [type(e).__name__ for e in events]
    assert names.index("MemoryDiffReady") < names.index("StrategyAnswerReady") < names.index("BuildComplete")


async def test_structured_outputs_are_stored_per_stage_and_session_is_persisted(rig):
    await rig.run()
    saved = rig.storage.get_session(rig.session.id)
    for key in ["deal_brief", "retrieval", "research", "personas", "strategy", "assets", "verifier"]:
        assert key in saved.artifacts
    assert set(saved.artifacts["assets"]) == {"deck"}
    assert [(s.stage_id, s.status) for s in saved.stages][-1] == ("strategy_answer", "done")
    assert saved.assets[0].asset_type == "deck"


async def test_stage_state_is_persisted_after_each_stage(rig):
    seen = []
    async for event in rig.runner.run(rig.session, rig.tenant):
        if isinstance(event, StageEvent) and event.stage.status == "done":
            on_disk = rig.storage.get_session(rig.session.id)
            seen.append((event.stage.stage_id, [s.stage_id for s in on_disk.stages if s.status == "done"]))
    assert seen[0] == ("deal_brief", ["deal_brief"])
    assert seen[3] == ("personas", ["deal_brief", "retrieval", "research", "personas"])


async def test_deal_brief_request_carries_transcript_board_memory_and_tenant_block(rig):
    rig.session.memory_snapshot = {"/accounts/globex/profile.md": "uses SAP"}
    await rig.run()
    request = rig.doubles.llm.calls("deal_brief")[0]
    assert "Memoria: Hi" in request.user and "Rep: Hello" in request.user
    assert "uses SAP" in request.user and "2026-10-23" in request.user
    assert "pain: [open] -" in request.user
    assert "Payroll software" in request.tenant_block
    entry = next(e for e in CHAIN if e["id"] == "deal_brief")
    assert (request.max_tokens, request.schema, request.effort) == (
        entry["max_tokens"], entry["schema"], entry["effort"])


async def test_retrieval_is_skipped_without_documents(tmp_path, doubles, profile, render_recorder):
    rig = Rig(tmp_path, doubles, profile)
    events = await rig.run()
    skipped = [e.stage for e in events if isinstance(e, StageEvent) and e.stage.status == "skipped"]
    assert [(s.stage_id, s.detail) for s in skipped] == [("retrieval", "no documents")]
    assert (("retrieval", "running") in stage_events(events))
    assert rig.doubles.llm.calls("retrieval") == []
    assert rig.session.phase == "complete"


async def test_documents_for_other_accounts_do_not_count(tmp_path, doubles, profile, render_recorder):
    rig = Rig(tmp_path, doubles, profile)
    rig.add_document("other.md", b"x", "text/markdown", scope="account", account_slug="initech")
    events = await rig.run()
    assert ("retrieval", "skipped") in stage_events(events)
    assert rig.doubles.llm.calls("retrieval") == []


async def test_retrieval_sends_citation_enabled_document_blocks(tmp_path, doubles, profile, render_recorder):
    rig = Rig(tmp_path, doubles, profile)
    pdf = b"%PDF-1.4 fake"
    rig.add_document("spec.pdf", pdf, "application/pdf")
    rig.add_document("notes.txt", "café notes".encode(), "text/plain", scope="account",
                     account_slug="globex")
    word = DocxDocument()
    word.add_paragraph("Security overview")
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "SOC2"
    table.rows[0].cells[1].text = "yes"
    buffer = io.BytesIO()
    word.save(buffer)
    rig.add_document("sec.docx", buffer.getvalue(),
                     "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    rig.add_document("other.md", b"hidden", "text/markdown", scope="account", account_slug="initech")

    await rig.run()

    blocks = {b["title"]: b for b in rig.doubles.llm.calls("retrieval")[0].documents}
    assert set(blocks) == {"spec.pdf", "notes.txt", "sec.docx"}
    assert all(b["type"] == "document" and b["citations"] == {"enabled": True} for b in blocks.values())
    assert blocks["spec.pdf"]["source"] == {
        "type": "base64", "media_type": "application/pdf", "data": base64.b64encode(pdf).decode()}
    assert blocks["notes.txt"]["source"] == {"type": "text", "media_type": "text/plain",
                                             "data": "café notes"}
    docx_text = blocks["sec.docx"]["source"]["data"]
    assert "Security overview" in docx_text and "SOC2 | yes" in docx_text


async def test_corrupt_docx_fails_the_retrieval_stage_with_the_filename(tmp_path, doubles, profile,
                                                                        render_recorder):
    rig = Rig(tmp_path, doubles, profile)
    rig.add_document("broken.docx", b"not a zip", "application/vnd.openxmlformats-officedocument."
                     "wordprocessingml.document")
    with pytest.raises(Exception, match="broken.docx"):
        await rig.run()
    assert [(s.stage_id, s.status) for s in rig.session.stages][-1] == ("retrieval", "failed")


async def test_evidence_ids_are_renumbered_by_code_across_retrieval_and_research(rig):
    rig.doubles.llm.script.update(
        retrieval={"entries": [evidence("E7", "alpha"), evidence("E7", "beta"),
                               evidence("E9", "gamma", conflict_note="conflicts with E7")]},
        research={
            "questions": [{"question": "q", "outcome": "answered",
                           "evidence_ids": ["E1", "E2", "E3", "E99"]}],
            "entries": [evidence("E40", "web one"), evidence("E41", "web two")],
        },
    )
    await rig.run()

    ledger = rig.session.artifacts["ledger"]
    assert [(e["id"], e["claim"]) for e in ledger] == [
        ("E1", "alpha"), ("E2", "beta"), ("E3", "gamma"), ("E4", "web one"), ("E5", "web two")]
    assert ledger[2]["conflict_note"] == "conflicts with E1"
    retrieval_request = rig.doubles.llm.calls("retrieval")[0]
    assert "E1" in retrieval_request.user
    research_request = rig.doubles.llm.calls("research")[0]
    assert "Number" not in research_request.user
    assert "E4" in research_request.finalize_user
    assert "alpha" in research_request.user
    # a question may cite earlier ledger entries or the new ones, never ids that exist nowhere
    assert rig.session.artifacts["research"]["questions"][0]["evidence_ids"] == ["E1", "E2", "E3"]
    # later stages see the renumbered ledger only
    for stage in ["personas", "strategy", "asset:deck", "verifier"]:
        text = rig.doubles.llm.calls(stage)[0].user
        assert "E5" in text and "E40" not in text and "E41" not in text


async def test_research_numbers_from_one_when_retrieval_was_skipped(tmp_path, doubles, profile,
                                                                    render_recorder):
    doubles.llm.script["research"] = {"questions": [], "entries": [evidence("E12", "web")]}
    rig = Rig(tmp_path, doubles, profile)
    await rig.run()
    assert [e["id"] for e in rig.session.artifacts["ledger"]] == ["E1"]
    request = doubles.llm.calls("research")[0]
    assert "E1" in request.finalize_user
    assert request.tools == ["web_search", "web_fetch"] and request.max_tool_calls == 20


async def test_strategy_gets_previous_strategy_from_the_memory_snapshot(rig):
    rig.session.memory_snapshot = {"/accounts/globex/strategy.md": "OLD PLAN: lead with price",
                                   "/accounts/initech/strategy.md": "OTHER ACCOUNT"}
    await rig.run()
    user = rig.doubles.llm.calls("strategy")[0].user
    assert "OLD PLAN: lead with price" in user and "OTHER ACCOUNT" not in user


async def test_strategy_marks_first_session_when_memory_has_no_strategy(rig):
    await rig.run()
    assert "(none: first session)" in rig.doubles.llm.calls("strategy")[0].user


async def test_only_requested_assets_in_the_plan_are_generated(rig):
    rig.doubles.llm.script.update(
        deal_brief=brief(["deck", "blog"]),
        strategy={"asset_plan": [plan_entry("deck"), plan_entry("whitepaper"), plan_entry("blog"),
                                 plan_entry("deck")]},
    )
    await rig.run()
    stages = sorted(r.stage for r in rig.doubles.llm.requests if r.stage.startswith("asset:"))
    assert stages == ["asset:blog", "asset:deck"]
    assert [a.asset_type for a in rig.session.assets] == ["deck", "blog"]


async def test_tenant_default_assets_apply_when_the_brief_requests_none(rig):
    rig.doubles.llm.script.update(
        deal_brief=brief([]),
        strategy={"asset_plan": [plan_entry("deck"), plan_entry("document")]},
    )
    await rig.run()
    assert [a.asset_type for a in rig.session.assets] == ["document"]


async def test_asset_requests_use_the_chain_system_schema_and_token_budget(rig):
    await rig.run()
    request = rig.doubles.llm.calls("asset:deck")[0]
    assert request.system == ASSET_SYSTEMS["deck"]
    assert request.schema == ASSET_SCHEMAS["deck"]
    assert request.max_tokens == 6000
    assert "deck goal" in request.user


async def test_assets_run_concurrently_but_never_more_than_three_at_once(rig):
    types = ["deck", "document", "blog", "email_sequence", "whitepaper"]
    active = peak = 0

    async def slow(req):
        nonlocal active, peak
        active += 1
        peak = max(peak, active)
        await asyncio.sleep(0.02)
        active -= 1
        return asset_content(req.stage.split(":")[1])

    rig.doubles.llm.script.update(
        deal_brief=brief(types), strategy={"asset_plan": [plan_entry(t) for t in types]},
        **{"asset:*": slow},
    )
    await rig.run()
    assert peak == 3
    assert [a.asset_type for a in rig.session.assets] == types


async def test_brand_rules_combine_voice_forbidden_claims_and_campaign_constraints(rig):
    await rig.run()
    for stage in ["asset:deck", "verifier"]:
        user = rig.doubles.llm.calls(stage)[0].user
        assert "warm and plain" in user
        assert "cheapest on the market" in user and "guaranteed ROI" in user
        assert "plain words" in user


async def test_revise_then_pass_renders_the_revised_asset(rig):
    rig.doubles.llm.script.update(
        verifier=[verdict("needs_revision"), verdict("pass")],
        reviser={"title": "Revised deck", "placeholders": ["[PLACEHOLDER: x]"]},
    )
    events = await rig.run()
    assert len(rig.doubles.llm.calls("reviser")) == 1
    assert len(rig.doubles.llm.calls("verifier")) == 2
    assert "needs_revision" in rig.doubles.llm.calls("reviser")[0].user
    ready = [e.asset for e in events if isinstance(e, AssetReady)]
    assert [(a.title, a.verdict, a.placeholders) for a in ready] == [
        ("Revised deck", "pass", ["[PLACEHOLDER: x]"])]
    assert rig.session.artifacts["assets"]["deck"]["title"] == "Revised deck"
    assert "Revised deck" in rig.doubles.llm.calls("verifier")[1].user


async def test_revised_asset_is_the_one_rendered(rig, render_recorder):
    rig.doubles.llm.script.update(verifier=[verdict("needs_revision"), verdict("pass")],
                                  reviser={"title": "Revised deck", "placeholders": []})
    await rig.run()
    assert [c["content"]["title"] for c in render_recorder.calls] == ["Revised deck"]


async def test_second_failure_ships_as_needs_human_review_without_a_second_revision(rig, render_recorder):
    rig.doubles.llm.script["verifier"] = verdict("needs_revision")
    events = await rig.run()
    assert len(rig.doubles.llm.calls("reviser")) == 1
    assert len(rig.doubles.llm.calls("verifier")) == 2
    ready = [e.asset for e in events if isinstance(e, AssetReady)]
    assert [a.verdict for a in ready] == ["needs_human_review"]
    assert len(render_recorder.calls) == 1 and ready[0].files
    assert isinstance(events[-1], BuildComplete)
    assert rig.session.assets[0].verdict == "needs_human_review"


async def test_asset_without_title_gets_a_label_from_its_type(rig):
    rig.doubles.llm.script.update(
        deal_brief=brief(["email_sequence"]), strategy={"asset_plan": [plan_entry("email_sequence")]})
    await rig.run()
    assert rig.session.assets[0].title == "Email sequence"


async def test_render_receives_ledger_directory_basename_and_url_prefix(rig, render_recorder):
    await rig.run()
    call = render_recorder.calls[0]
    assert call["asset_type"] == "deck" and call["basename"] == "deck"
    assert call["url_prefix"] == f"/api/files/{rig.session.id}"
    assert call["out_dir"] == rig.storage.rendered_dir(rig.session.id)
    assert call["ledger"] == rig.session.artifacts["ledger"]
    assert rig.session.assets[0].files[0].url == f"/api/files/{rig.session.id}/deck.md"


async def test_memory_diff_is_computed_around_the_memory_writer_and_emitted(rig):
    rig.doubles.memory.files.update({"/accounts/globex/objections.md": "O1 open", "/keep.md": "k"})
    rig.doubles.agents.writes = {"/accounts/globex/objections.md": "O1 answered", "/new.md": "n"}
    events = await rig.run()
    diff = next(e for e in events if isinstance(e, MemoryDiffReady))
    assert [(c.path, c.change) for c in diff.changes] == [
        ("/accounts/globex/objections.md", "modified"), ("/new.md", "added")]
    assert diff.changes[0].before == "O1 open" and diff.changes[0].after == "O1 answered"
    assert rig.session.memory_diff == diff.changes
    assert rig.storage.get_session(rig.session.id).memory_diff == diff.changes


async def test_agents_run_with_the_right_roles_access_and_store(rig):
    events = await rig.run()
    writer, briefer = rig.doubles.agents.runs
    assert (writer["role"], writer["access"], writer["memory_store_id"]) == (
        "memory_writer", "read_write", "store-1")
    assert (briefer["role"], briefer["access"], briefer["memory_store_id"]) == (
        "briefer", "read_only", "store-1")
    assert TRACK_TEST_QUESTION in briefer["user"]
    assert "globex" in writer["user"] and "dana" in writer["user"]
    assert "deck" in writer["user"]
    answer = next(e for e in events if isinstance(e, StrategyAnswerReady))
    assert answer.text == rig.session.strategy_answer == "Since our last session: nothing"


async def test_missing_memory_store_fails_the_memory_writer_stage(
    tmp_path, doubles, profile, render_recorder,
):
    rig = Rig(tmp_path, doubles, profile, with_store=False)
    with pytest.raises(Exception, match="no memory store"):
        await rig.run()
    assert (rig.session.stages[-1].stage_id, rig.session.stages[-1].status) == ("memory_writer", "failed")
    assert doubles.agents.runs == []


async def test_stage_failure_emits_failed_event_marks_session_failed_and_reraises(rig):
    rig.doubles.llm.script["personas"] = RuntimeError("model overloaded")
    seen = []
    with pytest.raises(RuntimeError, match="model overloaded"):
        async for event in rig.runner.run(rig.session, rig.tenant):
            seen.append(event)

    last = seen[-1]
    assert isinstance(last, StageEvent)
    assert (last.stage.stage_id, last.stage.status, last.stage.detail) == (
        "personas", "failed", "model overloaded")
    assert last.stage.duration_ms is not None
    assert not any(isinstance(e, BuildComplete) for e in seen)
    assert rig.doubles.llm.calls("strategy") == []
    assert rig.storage.get_session(rig.session.id).phase == "failed"


async def test_exception_without_message_still_reports_a_detail(rig):
    rig.doubles.llm.script["deal_brief"] = TimeoutError()
    seen = []
    with pytest.raises(asyncio.TimeoutError):
        async for event in rig.runner.run(rig.session, rig.tenant):
            seen.append(event)
    assert seen[-1].stage.detail == "TimeoutError"


async def test_a_failing_asset_cancels_its_siblings_and_fails_the_assets_stage(rig):
    started = asyncio.Event()
    finished = []

    async def slow(req):
        started.set()
        await asyncio.sleep(5)
        finished.append(req.stage)
        return asset_content("deck")

    def boom(req):
        raise ValueError("schema rejected")

    rig.doubles.llm.script.update(
        deal_brief=brief(["deck", "blog"]),
        strategy={"asset_plan": [plan_entry("deck"), plan_entry("blog")]},
        **{"asset:deck": slow, "asset:blog": boom},
    )
    with pytest.raises(ValueError, match="schema rejected"):
        await asyncio.wait_for(rig.run(), timeout=2)
    assert finished == []
    assert rig.session.stages[-1].stage_id == "assets" and rig.session.stages[-1].status == "failed"


async def test_missing_structured_output_is_a_stage_failure(rig):
    rig.doubles.llm.script["strategy"] = None
    with pytest.raises(Exception, match="strategy returned no structured output"):
        await rig.run()
    assert rig.session.phase == "failed"
