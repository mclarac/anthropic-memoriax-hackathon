import pytest
from pptx import Presentation

from memoria.renderers import render_asset

LEDGER = [
    {"id": f"E{n}", "claim": f"Claim {n}.", "source": f"Source {n}", "source_date": "2025-03-04"}
    for n in range(1, 13)
]
LEDGER[1]["source_date"] = None

DECK = {
    "title": "Acme <Pitch> & Co",
    "slides": [
        {
            "section": "Opening",
            "headline": "Setup in two days",
            "bullets": [{"text": "Fast **setup**.[E1][E2]"}, {"text": "Unknown fact.[E99]"}],
            "speaker_notes": "Say it plainly.[E3]",
            "visual": "Timeline graphic",
        },
        {
            "section": "Close",
            "headline": "Next steps",
            "bullets": [{"text": "Pilot with [PLACEHOLDER: customer]"}],
            "speaker_notes": "Ask for the pilot.",
            "visual": None,
        },
    ],
    "placeholders": ["[PLACEHOLDER: customer]"],
}


def render(tmp_path, deck=DECK):
    files = render_asset("deck", deck, LEDGER, tmp_path, "pitch", "/api/files/s")
    return files, Presentation(str(tmp_path / "pitch.pptx"))


def body_text(slide):
    return "\n".join(
        shape.text_frame.text for shape in slide.placeholders if shape.placeholder_format.idx == 1
    )


def test_single_pptx_file(tmp_path):
    files, _ = render(tmp_path)
    assert [(f.format, f.url) for f in files] == [("pptx", "/api/files/s/pitch.pptx")]
    assert files[0].size_bytes == (tmp_path / "pitch.pptx").stat().st_size


def test_slide_structure_and_titles(tmp_path):
    _, prs = render(tmp_path)
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles == ["Acme <Pitch> & Co", "Setup in two days", "Next steps", "Sources"]


def test_bullets_hide_known_markers_and_flag_missing(tmp_path):
    _, prs = render(tmp_path)
    text = body_text(prs.slides[1])
    assert text == "Fast setup.\nUnknown fact.[E99: missing evidence]"


def test_notes_have_speaker_notes_visual_and_evidence(tmp_path):
    _, prs = render(tmp_path)
    notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "Say it plainly." in notes
    assert "Visual: Timeline graphic" in notes
    assert "Section: Opening" in notes
    assert "E1: Claim 1. Source: Source 1 (2025-03-04)" in notes
    assert "E2: Claim 2. Source: Source 2 (undated)" in notes
    assert "E3: Claim 3." in notes
    assert "E99: missing evidence" in notes
    assert "Visual:" not in prs.slides[2].notes_slide.notes_text_frame.text


def test_placeholder_run_is_bold(tmp_path):
    _, prs = render(tmp_path)
    runs = [r for p in prs.slides[2].placeholders[1].text_frame.paragraphs for r in p.runs]
    flagged = [r for r in runs if "PLACEHOLDER" in r.text]
    assert flagged and all(r.font.bold for r in flagged)
    assert not [r for r in runs if "Pilot" in r.text and r.font.bold]


def test_sources_slide_lists_every_cited_id_once(tmp_path):
    _, prs = render(tmp_path)
    sources = body_text(prs.slides[-1])
    for n in (1, 2, 3):
        assert f"E{n}: Claim {n}. Source: Source {n}" in sources
    assert sources.count("E1:") == 1
    assert "E99: missing evidence" in sources
    assert "E4:" not in sources


def test_many_sources_spill_onto_continuation_slides(tmp_path):
    deck = {
        "title": "Big",
        "slides": [
            {"section": "s", "headline": "h", "speaker_notes": "",
             "bullets": [{"text": "x" + "".join(f"[E{n}]" for n in range(1, 13))}], "visual": None}
        ],
        "placeholders": [],
    }
    _, prs = render(tmp_path, deck)
    source_slides = [s for s in prs.slides if s.shapes.title.text.startswith("Sources")]
    assert len(source_slides) == 2
    combined = "\n".join(body_text(s) for s in source_slides)
    assert all(f"E{n}:" in combined for n in range(1, 13))


def test_deck_without_slides_or_citations(tmp_path):
    _, prs = render(tmp_path, {"title": "Empty", "slides": [], "placeholders": []})
    assert [s.shapes.title.text for s in prs.slides] == ["Empty", "Sources"]
    assert body_text(prs.slides[1]) == "No evidence cited."


@pytest.mark.parametrize("bad", ["../x", "a/b", ""])
def test_basename_must_be_a_plain_filename(tmp_path, bad):
    with pytest.raises(ValueError):
        render_asset("deck", DECK, LEDGER, tmp_path, bad, "/u")
