import re

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from memoria.renderers import render_asset

LEDGER = [
    {"id": "E4", "claim": "Setup takes two days.", "source": "Guide s2", "source_date": "2025-01-02"},
    {"id": "E9", "claim": "Costs fell 20%.", "source": "https://x.test", "source_date": None},
]

SECTIONS = [
    {
        "heading": "Why <now> & here",
        "body_markdown": (
            "Intro with **bold** and *soft* text.[E4][E9] Link [site](https://example.com).\n\n"
            "## Sub\n\n- first [PLACEHOLDER: ROI figure]\n- second [E99]\n\n"
            "1. one\n2. two\n\nTag <script>alert(1)</script> & done"
        ),
    },
    {"heading": "Second", "body_markdown": "Again.[E9]"},
]

DOCUMENT = {"kind": "one_pager", "title": "Acme & Co <Brief>", "sections": SECTIONS, "placeholders": []}
WHITEPAPER = {
    "title": "State of Setup",
    "executive_summary": "Summary says setup is quick.[E4]",
    "sections": SECTIONS,
    "placeholders": [],
}


def render(tmp_path, asset_type, content):
    files = render_asset(asset_type, content, LEDGER, tmp_path, "asset", "/api/files/s")
    return {f.format: f for f in files}


def pdf_lines(path):
    """Text shown on each line of an uncompressed reportlab PDF, split-string operators rejoined."""
    raw = path.read_bytes()
    chunks = re.split(rb"T\*|ET", raw)
    strings = [re.findall(rb"\(((?:\\.|[^\\)])*)\) Tj", chunk) for chunk in chunks]
    lines = [re.sub(r"\\(.)", r"\1", b"".join(found).decode("latin-1")) for found in strings]
    return [line for line in lines if line]


def paragraphs(path):
    return Document(str(path)).paragraphs


@pytest.mark.parametrize("asset_type,content", [("document", DOCUMENT), ("whitepaper", WHITEPAPER)])
def test_docx_and_pdf_pair_with_urls(tmp_path, asset_type, content):
    files = render(tmp_path, asset_type, content)
    assert list(files) == ["docx", "pdf"]
    assert files["docx"].url == "/api/files/s/asset.docx"
    assert files["pdf"].size_bytes == (tmp_path / "asset.pdf").stat().st_size


def test_docx_title_headings_and_levels(tmp_path):
    render(tmp_path, "document", DOCUMENT)
    paras = paragraphs(tmp_path / "asset.docx")
    assert (paras[0].style.name, paras[0].text) == ("Title", "Acme & Co <Brief>")
    headings = [(p.style.name, p.text) for p in paras if p.style.name.startswith("Heading")]
    assert headings == [
        ("Heading 1", "Why <now> & here"),
        ("Heading 3", "Sub"),
        ("Heading 1", "Second"),
        ("Heading 1", "References"),
    ]


def test_docx_formatting_citations_and_references(tmp_path):
    render(tmp_path, "document", DOCUMENT)
    paras = paragraphs(tmp_path / "asset.docx")
    intro = next(p for p in paras if p.text.startswith("Intro"))
    assert intro.text.startswith("Intro with bold and soft text.[1][2] Link site.")
    bold = [r.text for r in intro.runs if r.bold]
    italic = [r.text for r in intro.runs if r.italic]
    assert bold == ["bold"] and italic == ["soft"]
    assert [r.text for r in intro.runs if r.font.superscript] == ["[1]", "[2]"]
    refs = [p.text for p in paras[paras.index(next(p for p in paras if p.text == "References")) + 1:]]
    assert refs == [
        "[1] Setup takes two days. Source: Guide s2 (2025-01-02)",
        "[2] Costs fell 20%. Source: https://x.test (undated)",
    ]


def test_docx_hyperlink_placeholder_and_missing_evidence(tmp_path):
    render(tmp_path, "document", DOCUMENT)
    doc = Document(str(tmp_path / "asset.docx"))
    rels = [r.target_ref for r in doc.part.rels.values() if r.is_external]
    assert "https://example.com" in rels
    flagged = [
        r for p in doc.paragraphs for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW
    ]
    assert [r.text for r in flagged] == ["[PLACEHOLDER: ROI figure]"]
    assert any("[E99: missing evidence]" in p.text for p in doc.paragraphs)


def test_docx_lists_and_raw_markup_stays_literal(tmp_path):
    render(tmp_path, "document", DOCUMENT)
    texts = [(p.style.name, p.text) for p in paragraphs(tmp_path / "asset.docx")]
    assert ("List Bullet", "first [PLACEHOLDER: ROI figure]") in texts
    assert ("List Paragraph", "1. one") in texts and ("List Paragraph", "2. two") in texts
    assert any(t == "Tag <script>alert(1)</script> & done" for _, t in texts)


def test_whitepaper_puts_executive_summary_first(tmp_path):
    render(tmp_path, "whitepaper", WHITEPAPER)
    paras = paragraphs(tmp_path / "asset.docx")
    assert [p.text for p in paras[:3]] == [
        "State of Setup", "Executive summary", "Summary says setup is quick.[1]",
    ]
    assert paras[1].style.name == "Heading 1"


def test_pdf_has_header_title_and_escapes_markup(tmp_path):
    render(tmp_path, "document", DOCUMENT)
    assert (tmp_path / "asset.pdf").read_bytes().startswith(b"%PDF")
    lines = pdf_lines(tmp_path / "asset.pdf")
    assert "Acme & Co <Brief>" in lines
    assert "Tag <script>alert(1)</script> & done" in lines
    assert any("[E99: missing evidence]" in line for line in lines)
    assert "References" in lines
    assert "[1] Setup takes two days. Source: Guide s2 (2025-01-02)" in lines
    assert any("[PLACEHOLDER: ROI figure]" in line for line in lines)


def test_pdf_whitepaper_summary_precedes_sections(tmp_path):
    render(tmp_path, "whitepaper", WHITEPAPER)
    lines = pdf_lines(tmp_path / "asset.pdf")
    assert lines.index("Executive summary") < lines.index("Why <now> & here")


def test_empty_sections_still_render(tmp_path):
    content = {"kind": "one_pager", "title": "Bare", "sections": [], "placeholders": []}
    render(tmp_path, "document", content)
    assert (tmp_path / "asset.pdf").read_bytes().startswith(b"%PDF")
    assert paragraphs(tmp_path / "asset.docx")[0].text == "Bare"
