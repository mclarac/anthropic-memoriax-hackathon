"""The backfill lane: turns a finished interview into a verified campaign.

BuildRunner runs the stages of chain_prompts.CHAIN after the interview, in order, and
yields the wire events the UI renders. Every structured output is stored on the session
so a stage can be inspected after the fact, and the session is persisted after each stage.
"""

from __future__ import annotations

import asyncio
import base64
import io
import json
import logging
import re
import time
from collections.abc import AsyncIterator, Awaitable, Callable
from datetime import date
from functools import partial
from typing import TYPE_CHECKING, Any

from docx import Document as DocxDocument

from memoria.chain_prompts import (
    ASSET_USER,
    CHAIN,
    DEAL_BRIEF_USER,
    MEMORY_WRITER_USER,
    PERSONAS_USER,
    RESEARCH_FINALIZE,
    RESEARCH_USER,
    RETRIEVAL_USER,
    REVISER,
    REVISER_USER,
    STRATEGY_ANSWER_USER,
    STRATEGY_USER,
    TENANT_BLOCK,
    TRACK_TEST_QUESTION,
    VERIFIER_USER,
)
from memoria.models import (
    AssetReady,
    AssetRef,
    BuildComplete,
    DocumentRef,
    MemoryDiffReady,
    Session,
    StageEvent,
    StageId,
    StageState,
    StrategyAnswerReady,
    Tenant,
)
from memoria.ports import LLMRequest, diff_snapshots
from memoria.renderers import render_asset
from memoria.storage import Storage

if TYPE_CHECKING:
    from memoria.app import Ports

logger = logging.getLogger(__name__)

ASSET_CONCURRENCY = 3
NO_STRATEGY_YET = "(none: first session)"
NOTHING_YET = "(none)"

_CHAIN = {entry["id"]: entry for entry in CHAIN}
_EVIDENCE_REF = re.compile(r"\bE\d+\b")

Event = StageEvent | AssetReady | MemoryDiffReady | StrategyAnswerReady | BuildComplete
Emit = Callable[[Event], None]


class StageError(RuntimeError):
    """A stage could not produce a usable result."""


def _json(value: Any) -> str:
    return json.dumps(value, indent=2, ensure_ascii=False)


def _ledger_text(ledger: list[dict[str, Any]]) -> str:
    return _json(ledger) if ledger else NOTHING_YET


def _next_evidence_id(ledger: list[dict[str, Any]]) -> str:
    return f"E{len(ledger) + 1}"


def _renumber(
    entries: list[dict[str, Any]], ledger: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    """Assign sequential ids after the existing ledger; return the entries and the old-to-new id map.

    The model's own numbering is never trusted: it can repeat, skip or collide with ids
    already in the ledger, and every later stage cites by id.
    """
    mapping: dict[str, str] = {}
    for offset, entry in enumerate(entries, start=len(ledger) + 1):
        mapping.setdefault(entry["id"], f"E{offset}")
    renumbered = []
    for offset, entry in enumerate(entries, start=len(ledger) + 1):
        note = entry.get("conflict_note")
        if note:
            note = _EVIDENCE_REF.sub(lambda m: mapping.get(m.group(), m.group()), note)
        renumbered.append({**entry, "id": f"E{offset}", "conflict_note": note})
    return renumbered, mapping


def _docx_text(data: bytes, filename: str) -> str:
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:  # python-docx surfaces corrupt files as several unrelated types
        raise StageError(f"could not read {filename}: {exc}") from exc
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _document_block(ref: DocumentRef, data: bytes) -> dict[str, Any]:
    if ref.media_type == "application/pdf":
        source = {"type": "base64", "media_type": "application/pdf",
                  "data": base64.standard_b64encode(data).decode("ascii")}
    else:
        text = _docx_text(data, ref.filename) if ref.filename.lower().endswith(".docx") \
            else data.decode("utf-8", errors="replace")
        source = {"type": "text", "media_type": "text/plain", "data": text}
    return {"type": "document", "source": source, "title": ref.filename, "citations": {"enabled": True}}


def _format_memory(snapshot: dict[str, str]) -> str:
    if not snapshot:
        return "(empty: first session)"
    return "\n\n".join(f"--- {path}\n{content}" for path, content in sorted(snapshot.items()))


def _format_transcript(session: Session) -> str:
    return "\n".join(f"{'Memoria' if t.speaker == 'memoria' else 'Rep'}: {t.text}" for t in session.turns)


def _brand_rules(tenant: Tenant, brief: dict[str, Any]) -> str:
    campaign = brief["campaign"]
    forbidden = list(dict.fromkeys([*tenant.profile.forbidden_claims, *campaign["forbidden_claims"]]))
    sections = [
        f"Brand voice: {tenant.profile.brand_voice or '(not specified)'}",
        "Forbidden claims:\n" + ("\n".join(f"- {c}" for c in forbidden) if forbidden else NOTHING_YET),
        f"Campaign constraints: {campaign['brand_constraints'] or '(not specified)'}",
    ]
    return "\n\n".join(sections)


def _asset_title(asset_type: str, content: dict[str, Any]) -> str:
    return content.get("title") or asset_type.replace("_", " ").capitalize()


class BuildRunner:
    def __init__(self, ports: Ports, storage: Storage, today: Callable[[], date]) -> None:
        self._ports = ports
        self._storage = storage
        self._today = today

    async def run(self, session: Session, tenant: Tenant) -> AsyncIterator[Event]:
        session.phase = "building"
        session.stages = []
        self._storage.save_session(session)
        bodies: list[tuple[StageId, Callable[..., Awaitable[None]]]] = [
            ("deal_brief", self._deal_brief),
            ("retrieval", self._retrieval),
            ("research", self._research),
            ("personas", self._personas),
            ("strategy", self._strategy),
            ("assets", self._assets),
            ("verifier", self._verifier),
            ("memory_writer", self._memory_writer),
            ("strategy_answer", self._strategy_answer),
        ]
        for stage_id, body in bodies:
            async for event in self._stage(session, stage_id, partial(body, session, tenant)):
                yield event
        session.phase = "complete"
        self._storage.save_session(session)
        yield BuildComplete()

    async def _stage(
        self, session: Session, stage_id: StageId, body: Callable[[StageState, Emit], Awaitable[None]],
    ) -> AsyncIterator[Event]:
        """Run one stage body, forwarding the events it emits as they happen, then report its outcome."""
        state = StageState(stage_id=stage_id, status="running")
        session.stages.append(state)
        yield StageEvent(stage=state.model_copy())
        started = time.monotonic()
        queue: asyncio.Queue[Event | None] = asyncio.Queue()
        task = asyncio.create_task(body(state, queue.put_nowait))
        task.add_done_callback(lambda _: queue.put_nowait(None))
        try:
            while (event := await queue.get()) is not None:
                yield event
            task.result()
        except Exception as exc:
            state.status = "failed"
            state.detail = str(exc) or type(exc).__name__
            state.duration_ms = round((time.monotonic() - started) * 1000)
            session.phase = "failed"
            self._storage.save_session(session)
            yield StageEvent(stage=state.model_copy())
            raise
        finally:
            task.cancel()
        if state.status == "running":
            state.status = "done"
        state.duration_ms = round((time.monotonic() - started) * 1000)
        self._storage.save_session(session)
        yield StageEvent(stage=state.model_copy())

    async def _structured(
        self, stage: str, user: str, tenant: Tenant, *, system: str, schema: dict[str, Any],
        max_tokens: int, effort: str, **extra: Any,
    ) -> dict[str, Any]:
        request = LLMRequest(
            stage=stage, system=system, user=user, max_tokens=max_tokens, schema=schema, effort=effort,
            tenant_block=_tenant_block(tenant), **extra,
        )
        result = await self._ports.llm.complete(request)
        if result.data is None:
            raise StageError(f"{stage} returned no structured output")
        return result.data

    async def _chain_call(self, stage_id: str, user: str, tenant: Tenant, **extra: Any) -> dict[str, Any]:
        entry = _CHAIN[stage_id]
        return await self._structured(
            stage_id, user, tenant, system=entry["system"], schema=entry["schema"],
            max_tokens=entry["max_tokens"], effort=entry["effort"], **extra,
        )

    # Stage bodies

    async def _deal_brief(self, session: Session, tenant: Tenant, state: StageState, emit: Emit) -> None:
        user = DEAL_BRIEF_USER.substitute(
            memory_snapshot=_format_memory(session.memory_snapshot),
            slot_board=session.board.render(),
            transcript=_format_transcript(session),
            today=self._today().isoformat(),
        )
        session.artifacts["deal_brief"] = await self._chain_call("deal_brief", user, tenant)

    async def _retrieval(self, session: Session, tenant: Tenant, state: StageState, emit: Emit) -> None:
        refs = [
            d for d in tenant.documents
            if d.scope == "tenant" or d.account_slug == session.account_slug
        ]
        if not refs:
            state.status = "skipped"
            state.detail = "no documents"
            return
        blocks = [_document_block(ref, self._storage.read_document(tenant.id, ref)) for ref in refs]
        ledger = session.artifacts.setdefault("ledger", [])
        user = RETRIEVAL_USER.substitute(
            deal_brief=_json(session.artifacts["deal_brief"]),
            first_id=_next_evidence_id(ledger), today=self._today().isoformat(),
        )
        output = await self._chain_call("retrieval", user, tenant, documents=blocks)
        entries, _ = _renumber(output["entries"], ledger)
        ledger.extend(entries)
        session.artifacts["retrieval"] = {"entries": entries}

    async def _research(self, session: Session, tenant: Tenant, state: StageState, emit: Emit) -> None:
        entry = _CHAIN["research"]
        ledger = session.artifacts.setdefault("ledger", [])
        user = RESEARCH_USER.substitute(
            deal_brief=_json(session.artifacts["deal_brief"]),
            evidence_ledger=_ledger_text(ledger), today=self._today().isoformat(),
        )
        output = await self._chain_call(
            "research", user, tenant,
            tools=entry["tools"], max_tool_calls=entry["max_tool_calls"],
            finalize_user=RESEARCH_FINALIZE.substitute(first_id=_next_evidence_id(ledger)),
        )
        known = {e["id"] for e in ledger}
        entries, mapping = _renumber(output["entries"], ledger)
        questions = [
            {**q, "evidence_ids": [
                mapping.get(i, i) for i in q["evidence_ids"] if i in mapping or i in known
            ]}
            for q in output["questions"]
        ]
        ledger.extend(entries)
        session.artifacts["research"] = {"questions": questions, "entries": entries}

    async def _personas(self, session: Session, tenant: Tenant, state: StageState, emit: Emit) -> None:
        user = PERSONAS_USER.substitute(
            deal_brief=_json(session.artifacts["deal_brief"]),
            evidence_ledger=_ledger_text(session.artifacts["ledger"]),
        )
        session.artifacts["personas"] = await self._chain_call("personas", user, tenant)

    async def _strategy(self, session: Session, tenant: Tenant, state: StageState, emit: Emit) -> None:
        previous = session.memory_snapshot.get(f"/accounts/{session.account_slug}/strategy.md")
        user = STRATEGY_USER.substitute(
            previous_strategy=previous or NO_STRATEGY_YET,
            deal_brief=_json(session.artifacts["deal_brief"]),
            personas=_json(session.artifacts["personas"]),
            evidence_ledger=_ledger_text(session.artifacts["ledger"]),
            today=self._today().isoformat(),
        )
        session.artifacts["strategy"] = await self._chain_call("strategy", user, tenant)

    async def _assets(self, session: Session, tenant: Tenant, state: StageState, emit: Emit) -> None:
        brief, strategy = session.artifacts["deal_brief"], session.artifacts["strategy"]
        requested = brief["campaign"]["assets_requested"] or tenant.profile.default_assets
        plan: dict[str, dict[str, Any]] = {}
        for item in strategy["asset_plan"]:
            if item["asset_type"] in requested:
                plan.setdefault(item["asset_type"], item)

        entry = _CHAIN["assets"]
        shared = {
            "strategy": _json(strategy),
            "personas": _json(session.artifacts["personas"]),
            "evidence_ledger": _ledger_text(session.artifacts["ledger"]),
            "brand_rules": _brand_rules(tenant, brief),
        }
        limit = asyncio.Semaphore(ASSET_CONCURRENCY)

        async def generate(asset_type: str) -> dict[str, Any]:
            async with limit:
                user = ASSET_USER.substitute(asset_plan_entry=_json(plan[asset_type]), **shared)
                return await self._structured(
                    f"asset:{asset_type}", user, tenant, system=entry["systems"][asset_type],
                    schema=entry["schemas"][asset_type], max_tokens=entry["max_tokens"][asset_type],
                    effort=entry["effort"],
                )

        tasks = {asset_type: asyncio.create_task(generate(asset_type)) for asset_type in plan}
        try:
            contents = await asyncio.gather(*tasks.values())
        finally:
            for task in tasks.values():
                task.cancel()
        session.artifacts["assets"] = dict(zip(plan, contents, strict=True))

    async def _verifier(self, session: Session, tenant: Tenant, state: StageState, emit: Emit) -> None:
        assets: dict[str, dict[str, Any]] = session.artifacts["assets"]
        reports: dict[str, dict[str, Any]] = {}
        session.artifacts["verifier"] = reports
        ledger = session.artifacts["ledger"]
        shared = {
            "evidence_ledger": _ledger_text(ledger),
            "brand_rules": _brand_rules(tenant, session.artifacts["deal_brief"]),
        }
        limit = asyncio.Semaphore(ASSET_CONCURRENCY)

        async def verify(asset_type: str) -> dict[str, Any]:
            user = VERIFIER_USER.substitute(
                asset_type=asset_type, asset_json=_json(assets[asset_type]), **shared,
            )
            return await self._chain_call("verifier", user, tenant)

        async def revise(asset_type: str, report: dict[str, Any]) -> dict[str, Any]:
            user = REVISER_USER.substitute(
                asset_json=_json(assets[asset_type]), verifier_report=_json(report),
            )
            return await self._structured(
                "reviser", user, tenant, system=REVISER["system"], schema=REVISER["schemas"][asset_type],
                max_tokens=_CHAIN["assets"]["max_tokens"][asset_type], effort=REVISER["effort"],
            )

        async def check(asset_type: str) -> AssetRef:
            async with limit:
                report = await verify(asset_type)
                if report["asset_verdict"] != "pass":
                    assets[asset_type] = await revise(asset_type, report)
                    report = await verify(asset_type)
                reports[asset_type] = report
                content = assets[asset_type]
                files = await asyncio.to_thread(
                    render_asset, asset_type, content, ledger,
                    self._storage.rendered_dir(session.id), asset_type, f"/api/files/{session.id}",
                )
                return AssetRef(
                    asset_type=asset_type, title=_asset_title(asset_type, content),
                    verdict="pass" if report["asset_verdict"] == "pass" else "needs_human_review",
                    placeholders=content.get("placeholders", []), files=files,
                )

        tasks = [asyncio.create_task(check(asset_type)) for asset_type in assets]
        refs: dict[str, AssetRef] = {}
        try:
            for finished in asyncio.as_completed(tasks):
                ref = await finished
                refs[ref.asset_type] = ref
                emit(AssetReady(asset=ref))
        finally:
            for task in tasks:
                task.cancel()
        session.assets = [refs[asset_type] for asset_type in assets]

    async def _memory_writer(self, session: Session, tenant: Tenant, state: StageState, emit: Emit) -> None:
        store_id = self._require_memory_store(tenant)
        summary = "\n".join(
            f"- {a.asset_type}: {a.title} ({a.verdict})"
            + (f"; placeholders: {', '.join(a.placeholders)}" if a.placeholders else "")
            for a in session.assets
        ) or NOTHING_YET
        user = MEMORY_WRITER_USER.substitute(
            account_slug=session.account_slug, rep_slug=session.rep_slug, today=self._today().isoformat(),
            deal_brief=_json(session.artifacts["deal_brief"]),
            strategy=_json(session.artifacts["strategy"]), asset_summary=summary,
        )
        before = await self._ports.memory.snapshot(store_id)
        await self._ports.agents.run(
            role="memory_writer", memory_store_id=store_id, access="read_write",
            user=user, title=f"memory update: {session.account_slug} ({session.id})",
        )
        after = await self._ports.memory.snapshot(store_id)
        session.memory_diff = diff_snapshots(before, after)
        emit(MemoryDiffReady(changes=session.memory_diff))

    async def _strategy_answer(
        self, session: Session, tenant: Tenant, state: StageState, emit: Emit,
    ) -> None:
        user = STRATEGY_ANSWER_USER.substitute(
            account_slug=session.account_slug, today=self._today().isoformat(),
            strategy=_json(session.artifacts["strategy"]), question=TRACK_TEST_QUESTION,
        )
        result = await self._ports.agents.run(
            role="briefer", memory_store_id=self._require_memory_store(tenant), access="read_only",
            user=user, title=f"strategy answer: {session.account_slug} ({session.id})",
        )
        session.strategy_answer = result.text
        emit(StrategyAnswerReady(text=result.text))

    @staticmethod
    def _require_memory_store(tenant: Tenant) -> str:
        if tenant.memory_store_id is None:
            raise StageError(f"tenant {tenant.id} has no memory store")
        return tenant.memory_store_id


def _tenant_block(tenant: Tenant) -> str:
    return TENANT_BLOCK.substitute(tenant_profile=tenant.profile.model_dump_json(indent=2))
