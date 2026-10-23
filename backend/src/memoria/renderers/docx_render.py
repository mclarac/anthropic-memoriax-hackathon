"""Word output via python-docx."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph

from memoria.renderers.blocks import Heading, ListBlock
from memoria.renderers.inline import Run
from memoria.renderers.longform import LongForm

_MAX_HEADING = 9
_LINK_BLUE = RGBColor(0x0B, 0x5C, 0xAD)
_MISSING_RED = RGBColor(0xB0, 0x00, 0x20)


def render_docx(longform: LongForm, path: Path) -> None:
    document = Document()
    document.core_properties.title = longform.title
    document.add_heading(longform.title, 0)
    for block in longform.blocks:
        if isinstance(block, Heading):
            _add_runs(document.add_heading("", min(block.level, _MAX_HEADING)), block.runs)
        elif isinstance(block, ListBlock):
            _add_list(document, block)
        else:
            _add_runs(document.add_paragraph(), block.runs)
    document.save(str(path))


def _add_list(document, block: ListBlock) -> None:
    for position, item in enumerate(block.items, start=1):
        if block.ordered:
            # The built-in numbered style continues numbering across lists, so numbers are written out.
            paragraph = document.add_paragraph(style="List Paragraph")
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.add_run(f"{position}. ")
        else:
            paragraph = document.add_paragraph(style="List Bullet")
        _add_runs(paragraph, item)


def _add_runs(paragraph: DocxParagraph, runs: list[Run]) -> None:
    for run in runs:
        if run.href:
            _add_link(paragraph, run)
        else:
            _style(paragraph.add_run(_display(run)), run)


def _display(run: Run) -> str:
    return f"[{run.text}]" if run.kind == "cite" else run.text


def _style(docx_run, run: Run) -> None:
    docx_run.bold = run.bold or run.kind in {"placeholder", "missing"}
    docx_run.italic = run.italic
    if run.kind == "cite":
        docx_run.font.superscript = True
    elif run.kind == "placeholder":
        docx_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    elif run.kind == "missing":
        docx_run.font.color.rgb = _MISSING_RED


def _add_link(paragraph: DocxParagraph, run: Run) -> None:
    relationship = paragraph.part.relate_to(run.href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    docx_run = paragraph.add_run(_display(run))
    _style(docx_run, run)
    docx_run.font.underline = True
    docx_run.font.color.rgb = _LINK_BLUE
    hyperlink.append(docx_run._r)
    paragraph._p.append(hyperlink)
