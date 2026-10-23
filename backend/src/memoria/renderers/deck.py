"""Slide deck output via python-pptx."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from memoria.renderers.citations import Citations
from memoria.renderers.inline import MARKER, Run, parse_inline

_TITLE_LAYOUT, _CONTENT_LAYOUT = 0, 1
_SOURCES_PER_SLIDE = 7
_MISSING_RED = RGBColor(0xB0, 0x00, 0x20)
_MARGIN = Inches(0.6)


def render_deck(content: dict, ledger: list[dict], out_dir: Path, basename: str) -> list[Path]:
    citations = Citations(ledger)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    _add_title_slide(presentation, content["title"])
    cited: dict[str, None] = {}
    for slide in content["slides"]:
        texts = [bullet["text"] for bullet in slide["bullets"]] + [slide["speaker_notes"]]
        slide_ids = _marker_ids(texts)
        cited.update(dict.fromkeys(slide_ids))
        bullets = [citations.drop_known(parse_inline(bullet["text"])) for bullet in slide["bullets"]]
        notes = _notes(slide, slide_ids, citations)
        _add_content_slide(presentation, slide["headline"], bullets, notes)
    _add_sources(presentation, list(cited), citations)
    path = out_dir / f"{basename}.pptx"
    presentation.save(str(path))
    return [path]


def _marker_ids(texts: list[str]) -> list[str]:
    ids: dict[str, None] = {}
    for text in texts:
        for match in MARKER.finditer(text):
            ids.setdefault(f"E{match[1]}")
    return list(ids)


def _evidence_line(evidence_id: str, citations: Citations) -> str:
    if not citations.has(evidence_id):
        return f"{evidence_id}: missing evidence"
    return f"{evidence_id}: {citations.reference_text(citations.entry(evidence_id))}"


def _notes(slide: dict, ids: list[str], citations: Citations) -> str:
    parts = [f"Section: {slide['section']}", slide["speaker_notes"]]
    if slide.get("visual"):
        parts.append(f"Visual: {slide['visual']}")
    if ids:
        parts.append("Evidence:\n" + "\n".join(_evidence_line(i, citations) for i in ids))
    return "\n\n".join(part for part in parts if part)


def _drop_placeholder(shape) -> None:
    shape._element.getparent().remove(shape._element)


def _place(shape, left, top, width, height) -> None:
    shape.left, shape.top, shape.width, shape.height = left, top, width, height


def _add_title_slide(presentation: Presentation, title: str) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[_TITLE_LAYOUT])
    _drop_placeholder(slide.placeholders[1])
    shape = slide.shapes.title
    _place(shape, Inches(1), Inches(2.5), presentation.slide_width - Inches(2), Inches(2))
    shape.text_frame.text = title
    shape.text_frame.paragraphs[0].font.size = Pt(44)


def _content_slide(presentation: Presentation, title: str):
    slide = presentation.slides.add_slide(presentation.slide_layouts[_CONTENT_LAYOUT])
    width = presentation.slide_width - 2 * _MARGIN
    heading = slide.shapes.title
    _place(heading, _MARGIN, Inches(0.4), width, Inches(1.2))
    heading.text_frame.text = title
    heading.text_frame.paragraphs[0].font.size = Pt(32)
    body = slide.placeholders[1]
    _place(body, _MARGIN, Inches(1.8), width, presentation.slide_height - Inches(2.3))
    return slide, body


def _add_content_slide(presentation, headline: str, bullets: list[list[Run]], notes: str) -> None:
    slide, body = _content_slide(presentation, headline)
    if bullets:
        _fill(body, bullets)
    else:
        _drop_placeholder(body)
    slide.notes_slide.notes_text_frame.text = notes


def _fill(body, bullets: list[list[Run]]) -> None:
    characters = sum(len(run.text) for runs in bullets for run in runs)
    size = Pt(24 if characters < 350 and len(bullets) <= 5 else 18)
    frame = body.text_frame
    frame.word_wrap = True
    for index, runs in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        for run in runs:
            pptx_run = paragraph.add_run()
            pptx_run.text = run.text
            pptx_run.font.size = size
            pptx_run.font.bold = run.bold or run.kind in {"placeholder", "missing"}
            pptx_run.font.italic = run.italic
            if run.kind == "missing":
                pptx_run.font.color.rgb = _MISSING_RED
            if run.href:
                pptx_run.hyperlink.address = run.href


def _add_sources(presentation: Presentation, ids: list[str], citations: Citations) -> None:
    if not ids:
        _sources_slide(presentation, "Sources", ["No evidence cited."])
        return
    pages = [ids[i:i + _SOURCES_PER_SLIDE] for i in range(0, len(ids), _SOURCES_PER_SLIDE)]
    for number, page in enumerate(pages, start=1):
        title = "Sources" if number == 1 else f"Sources ({number})"
        _sources_slide(presentation, title, [_evidence_line(i, citations) for i in page])


def _sources_slide(presentation: Presentation, title: str, lines: list[str]) -> None:
    slide, body = _content_slide(presentation, title)
    _fill(body, [[Run(line)] for line in lines])
    for paragraph in body.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(16)
